"""Pick the documents of a downloaded tender that hold its rules, and join them into one PDF for the eligibility run.

A Tender Tiger download has many files: the RFP / NIT, a separate eligibility or scope document, corrigenda, but also BOQs,
price bids, drawings and blank forms. Each PDF / DOCX / TXT is scored on its file name and the text of its first two pages:

    analyse   score >= 5, or any corrigendum / addendum
    unsure    score 1..4 and under 30 pages: analysed too, and named as unsure on the Teams card
    skip      the rest (and every image, spreadsheet, zip, empty or duplicate file)

When nothing reaches 5, the largest PDF is taken ("main document guessed"). Deterministic: no AI is used to choose.
"""
import hashlib
import io
import os
import re
from typing import Any, Dict, List, Optional, Tuple

_TEXT_EXT = {".pdf", ".docx", ".txt"}
_MAX_FILE = 60 * 1024 * 1024
_UNSURE_MAX_PAGES = 30
MAX_PAGES = 400                 # the joined PDF stays below this: unsure files go first, then the smallest main ones

_NAME_RULES: List[Tuple[re.Pattern, int, str]] = [
    (re.compile(r"\b(rfp|rfq|nit|nib|eoi|itb|tender[\s_-]*doc|bid[\s_-]*doc|tender[\s_-]*document|bidding[\s_-]*doc)", re.I), 5, "name says RFP / NIT / bid document"),
    (re.compile(r"(eligib|pqc|pre[\s_-]*qualif|qualif)", re.I), 5, "name says eligibility / qualification"),
    (re.compile(r"(scope|technical[\s_-]*spec|spec|sow|terms[\s_-]*of[\s_-]*reference|\btor\b)", re.I), 3, "name says scope / specification"),
    (re.compile(r"(boq|bill[\s_-]*of[\s_-]*quant|price[\s_-]*bid|financial[\s_-]*bid|commercial[\s_-]*bid|price[\s_-]*schedule|drawing|"
                r"format|undertaking|affidavit|declaration|letter[\s_-]*of|\bform\b)", re.I), -5, "name says BOQ / price bid / drawing / form"),
]
_CORRIGENDUM = re.compile(r"(corrigend|addend|amendment|clarification|pre[\s_-]*bid[\s_-]*(reply|response|minutes))", re.I)
_TEXT_RULES: List[Tuple[re.Pattern, int, str]] = [
    (re.compile(r"notice\s+inviting\s+(e-?)?tender|request\s+for\s+(proposal|quotation)|instructions?\s+to\s+bidders|invitation\s+for\s+bids?|"
                r"expression\s+of\s+interest|tender\s+document", re.I), 5, "first pages: notice inviting tender / RFP"),
    (re.compile(r"eligibility\s+criteria|pre-?\s*qualification|qualifying\s+requirements?|qualification\s+criteria", re.I), 6,
     "first pages: eligibility criteria"),
    (re.compile(r"scope\s+of\s+work|technical\s+specifications?|terms\s+of\s+reference", re.I), 3, "first pages: scope / specifications"),
    (re.compile(r"bill\s+of\s+quantit|price\s+schedule|financial\s+bid|price\s+bid", re.I), -5, "first pages: BOQ / price bid"),
]


def _first_pages(name: str, data: bytes) -> Tuple[str, int]:
    """Text of the first two pages (or first 4000 characters) and the page count; ('', 0) when unreadable."""
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf":
            import fitz
            with fitz.open(stream=data, filetype="pdf") as doc:
                return "\n".join(doc[i].get_text() for i in range(min(2, doc.page_count))), doc.page_count
        if ext == ".docx":
            from docx import Document
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
            return text[:4000], max(1, len(text) // 3000)
        text = data.decode("utf-8", "replace")
        return text[:4000], max(1, len(text) // 3000)
    except Exception:
        return "", 0


def pick(files: List[Dict[str, Any]]) -> Dict[str, Any]:
    """files = [{file_name, content}] -> {"analyse": [...], "skipped": [...], "guessed": bool}; each item has name, reason, score."""
    analyse: List[Dict[str, Any]] = []
    skipped: List[Dict[str, Any]] = []
    seen: Dict[str, str] = {}
    candidates: List[Dict[str, Any]] = []
    for f in files:
        name, data = f["file_name"], f.get("content") or b""
        ext = os.path.splitext(name)[1].lower()
        if ext not in _TEXT_EXT:
            skipped.append({"name": name, "reason": f"{ext or 'no extension'} file (not a document with text)"})
            continue
        if not data:
            skipped.append({"name": name, "reason": "empty or could not be read"})
            continue
        if len(data) > _MAX_FILE:
            skipped.append({"name": name, "reason": "larger than 60 MB"})
            continue
        digest = hashlib.sha256(data).hexdigest()
        if digest in seen:
            skipped.append({"name": name, "reason": f"same file as {seen[digest]}"})
            continue
        seen[digest] = name
        text, pages = _first_pages(name, data)
        # the same document downloaded again under another name (Tender Tiger often lists one RFP several times)
        same = f"{pages}:" + re.sub(r"\s+", " ", text).strip()[:1500].lower()
        if len(same) > 200 and same in seen:
            skipped.append({"name": name, "reason": f"same document as {seen[same]}"})
            continue
        seen[same] = name
        score, why = 0, []
        for rx, pts, label in _NAME_RULES:
            if rx.search(name):
                score += pts
                why.append(label)
        for rx, pts, label in _TEXT_RULES:
            if rx.search(text):
                score += pts
                why.append(label)
        if pages > 10:
            score += 1
        elif 0 < pages <= 2:
            score -= 2
        corrigendum = bool(_CORRIGENDUM.search(name) or _CORRIGENDUM.search(text[:600]))
        candidates.append({"name": name, "content": data, "score": score, "pages": pages, "corrigendum": corrigendum,
                           "reason": "; ".join(why) or "no tender words found", "pdf": ext == ".pdf"})

    for c in candidates:
        if c["corrigendum"]:
            analyse.append({**c, "role": "corrigendum"})
        elif c["score"] >= 5:
            analyse.append({**c, "role": "main"})
        elif c["score"] >= 1 and c["reason"] != "no tender words found" and 0 < c["pages"] < _UNSURE_MAX_PAGES:
            analyse.append({**c, "role": "unsure"})
        else:
            skipped.append({"name": c["name"], "reason": f"score {c['score']}: {c['reason']}"})

    guessed = False
    if not any(a["role"] == "main" for a in analyse):
        pdfs = [c for c in candidates if c["pdf"] and not c["corrigendum"] and c not in analyse]
        pdfs += [a for a in analyse if a["role"] == "unsure" and a["pdf"]]
        if pdfs:
            best = max(pdfs, key=lambda c: (c["pages"], c["score"]))
            analyse = [a for a in analyse if a["name"] != best["name"]] + [{**best, "role": "main"}]
            skipped = [s for s in skipped if s["name"] != best["name"]]
            guessed = True
    # main documents first (most pages first), then unsure ones, corrigenda last so their changes come after the original wording
    order = {"main": 0, "unsure": 1, "corrigendum": 2}
    analyse.sort(key=lambda a: (order[a["role"]], -a["pages"]))
    while sum(a["pages"] for a in analyse) > MAX_PAGES and len(analyse) > 1:
        drop = next((a for a in analyse if a["role"] == "unsure"), None) or \
            min((a for a in analyse if a["role"] == "main"), key=lambda a: a["pages"], default=None)
        if drop is None or (drop["role"] == "main" and sum(1 for a in analyse if a["role"] == "main") == 1):
            break
        analyse.remove(drop)
        skipped.append({"name": drop["name"], "reason": f"left out to keep the analysis under {MAX_PAGES} pages"})
    return {"analyse": analyse, "skipped": skipped, "guessed": guessed}


def combine(docs: List[Dict[str, Any]]) -> Tuple[bytes, List[Dict[str, Any]]]:
    """One PDF made of every picked document, and the page map [{name, role, first, last}] (1-based pages of the joined PDF)."""
    import fitz
    out = fitz.open()
    page_map: List[Dict[str, Any]] = []
    for d in docs:
        first = out.page_count + 1
        if d["name"].lower().endswith(".pdf"):
            with fitz.open(stream=d["content"], filetype="pdf") as src:
                out.insert_pdf(src)
        else:
            text = _docx_text(d["content"]) if d["name"].lower().endswith(".docx") else d["content"].decode("utf-8", "replace")
            _text_pages(out, text)
        if out.page_count >= first:
            page_map.append({"name": d["name"], "role": d["role"], "first": first, "last": out.page_count})
    data = out.tobytes(garbage=3, deflate=True)
    out.close()
    return data, page_map


def _docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _text_pages(out, text: str, lines_per_page: int = 60) -> None:
    """Write plain text onto A4 pages (DOCX / TXT files are joined to the PDF this way)."""
    lines: List[str] = []
    for raw in text.splitlines() or [""]:
        while len(raw) > 110:
            cut = raw.rfind(" ", 0, 110)
            cut = cut if cut > 40 else 110
            lines.append(raw[:cut])
            raw = raw[cut:].lstrip()
        lines.append(raw)
    for i in range(0, len(lines), lines_per_page):
        page = out.new_page(width=595, height=842)
        page.insert_textbox((36, 36, 559, 806), "\n".join(lines[i:i + lines_per_page]), fontsize=8.5, fontname="helv")


def where(page: Optional[int], page_map: List[Dict[str, Any]]) -> str:
    """'Eligibility.pdf · p.4' for a page of the joined PDF."""
    if not page:
        return ""
    for m in page_map or []:
        if m["first"] <= page <= m["last"]:
            return f"{m['name']} · p.{page - m['first'] + 1}"
    return f"p.{page}"
