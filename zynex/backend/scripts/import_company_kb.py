"""Load the company knowledge-base workbook (Xenovex_Technology-KB.xlsx) into the zynex database.

Usage (from backend/):
    python -m scripts.import_company_kb ../Xenovex_Technology-KB.xlsx                 # dry run: reads, checks, prints, writes nothing
    python -m scripts.import_company_kb ../Xenovex_Technology-KB.xlsx --apply         # upsert by ID (safe to re-run)
    python -m scripts.import_company_kb ../Xenovex_Technology-KB.xlsx --apply --replace-demo   # also delete rows without an ID (the demo seed)

What is loaded            : sheets 01-07 (profile, financials, certificates, projects, employees, clients, technologies).
What is deliberately NOT  : sheet 00 (cover), sheet 08 (expected verdicts: that is an answer key, not company data) and the
                            "derived metrics" rows of sheet 02 (the engine computes averages itself).
Money                     : the workbook is in INR crore, the database in rupees, so values are multiplied by 1e7.
"""
import argparse
import asyncio
import os
import re
import sys
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl  # noqa: E402
from rapidfuzz import fuzz  # noqa: E402
from sqlalchemy import delete, select  # noqa: E402

from app.core.database import AsyncSessionLocal  # noqa: E402
from app.models.knowledge_base import Certification, Client, CompanyProfile, Employee, Project, Technology  # noqa: E402
from app.models.project_documents import FinancialRecord  # noqa: E402

CRORE = 1e7
SENSITIVE_KEYS = {"bank_account_number", "ifsc_code"}      # kept in the database, never put in a search index
_PREFIX = {"clients": "CLT-", "projects": "PRJ-", "employees": "EMP-", "certifications": "CERT-"}


def _rows(wb, sheet: str, id_prefix: str) -> List[tuple]:
    pattern = re.compile(r"^FY \d{4}") if id_prefix == "FY " else re.compile("^" + re.escape(id_prefix))
    return [r for r in wb[sheet].iter_rows(values_only=True) if r[0] and pattern.match(str(r[0]))]


def _date(v: Any) -> Optional[date]:
    if v is None or str(v).strip() in ("", "-", "N/A"):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d-%m-%Y", "%d %B %Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    return None


def _num(v: Any) -> Optional[float]:
    try:
        return float(str(v).replace(",", "").replace("₹", "").strip())
    except (TypeError, ValueError):
        return None


def _list(v: Any) -> List[str]:
    return [x.strip() for x in re.split(r"[,;]", str(v or "")) if x.strip()]


def _fy_end(label: str) -> Optional[int]:
    """'FY 2023-24' -> 2024 (the engine keys financial years by the year they end)."""
    m = re.search(r"(\d{4})\s*-\s*(\d{2,4})", str(label))
    if not m:
        return None
    end = m.group(2)
    return int(end) if len(end) == 4 else int(m.group(1)[:2] + end)


def read_workbook(path: str) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: Dict[str, Any] = {"warnings": []}

    # ---- profile: Attribute | Value | Note
    profile = {}
    for r in wb["01_Company_Profile"].iter_rows(min_row=5, values_only=True):
        if r[0] and r[1]:
            key = re.sub(r"[^a-z0-9]+", "_", str(r[0]).lower()).strip("_")
            profile[key] = (str(r[1]).strip(), str(r[2] or "").strip(), key in SENSITIVE_KEYS)
    out["profile"] = profile

    # ---- financial records: only rows that start with 'FY '
    fin = []
    for r in _rows(wb, "02_Financial_Records", "FY "):
        fy = _fy_end(r[0])
        if fy is None or _num(r[1]) is None:
            out["warnings"].append(f"financials: cannot read row {r[0]!r}")
            continue
        fin.append(dict(fiscal_year=fy, annual_turnover=_num(r[1]) * CRORE, net_profit=(_num(r[2]) or 0) * CRORE if _num(r[2]) is not None else None,
                        gross_assets=_num(r[3]) and _num(r[3]) * CRORE, total_liabilities=_num(r[4]) and _num(r[4]) * CRORE,
                        net_worth=_num(r[5]) and _num(r[5]) * CRORE, working_capital=_num(r[6]) and _num(r[6]) * CRORE,
                        ebitda=_num(r[7]) and _num(r[7]) * CRORE, current_ratio=_num(r[8]), audit_firm=r[9], audit_signoff_date=_date(r[10]),
                        notes=f"Imported from workbook ({r[0]})"))
    out["financials"] = fin

    certs = []
    for r in _rows(wb, "03_Certifications", "CERT-"):
        certs.append(dict(cert_code=r[0], name=str(r[1]).strip(), issuing_body=r[2], certificate_number=r[3], issued_on=_date(r[4]),
                          valid_until_date=_date(r[5]), valid_until=(_date(r[5]).isoformat() if _date(r[5]) else None), status=r[6]))
    out["certifications"] = certs

    projects = []
    for r in _rows(wb, "04_Past_Projects", "PRJ-"):
        months = _num(r[6])
        projects.append(dict(project_code=r[0], name=str(r[1]).strip(), client_name=str(r[2] or "").strip(),
                             total_value=(_num(r[3]) or 0) * CRORE if _num(r[3]) is not None else None,
                             start_date=_date(r[4]), completion_date=_date(r[5]),
                             duration_weeks=round(months * 4.345) if months else None,
                             status=str(r[7] or "").strip().lower() or "completed", domain=r[8], technologies=_list(r[9]),
                             order_no=r[10], description=str(r[12] or "").strip()))
    out["projects"] = projects

    emps = []
    for r in _rows(wb, "05_Employees_Key_Personnel", "EMP-"):
        emps.append(dict(employee_code=r[0], name=str(r[1]).strip(), role=str(r[2]).strip(), level=r[3], department=r[4],
                         experience_years=_num(r[5]), qualification=r[6], professional_certifications=r[7], skills=_list(r[8]),
                         is_active=str(r[9] or "").strip().lower() == "active"))
    out["employees"] = emps

    clients = []
    for r in _rows(wb, "06_Clients_Directory", "CLT-"):
        clients.append(dict(client_code=r[0], name=str(r[1]).strip(), client_type=r[2], industry=r[2], country=r[3], state_region=r[4],
                            engagement_status=r[5], contact_email=r[6], notes=str(r[7] or "").strip()))
    out["clients"] = clients

    techs = []
    for r in wb["07_Technology_Taxonomy"].iter_rows(min_row=5, values_only=True):
        if r[0] and r[1]:
            techs.append(dict(name=str(r[0]).strip(), canonical_id=str(r[1]).strip(), category=r[2], proficiency=r[3], experience_years=_num(r[4])))
    out["technologies"] = techs
    out["warnings"] += quality_checks(out)
    return out


def quality_checks(d: Dict[str, Any]) -> List[str]:
    """Things a person should fix in the workbook; the import still goes ahead."""
    w: List[str] = []
    today = date.today()
    years = sorted(f["fiscal_year"] for f in d["financials"])
    last_done = today.year if today.month > 3 else today.year - 1        # latest FY ended on 31 March
    if years:
        missing = [y for y in range(years[-1] + 1, last_done + 1)]
        if missing:
            w.append("financials end at FY " + f"{years[-1] - 1}-{str(years[-1])[2:]}; missing " +
                     ", ".join(f"FY {y - 1}-{str(y)[2:]}" for y in missing) + " (tenders asking for 'last 3 years' cannot be settled)")
    for c in d["certifications"]:
        if c["valid_until_date"] and c["valid_until_date"] <= today:
            w.append(f"{c['name'][:40]} EXPIRED on {c['valid_until_date']} but is marked {c['status']}")
        elif c["valid_until_date"] and c["valid_until_date"] <= today + timedelta(days=90):
            w.append(f"{c['name'][:40]} expires {c['valid_until_date']} (within 90 days)")
    # the same certificate carrying two different dates in the profile sheet and the certificates sheet
    for c in d["certifications"]:
        for key, (val, note, _) in d["profile"].items():
            tag = re.match(r"ISO[\s/]*(?:IEC\s*)?(\d{4,5})", c["name"])
            if tag and tag.group(1) in key.replace("_", "") and c["valid_until_date"]:
                m = re.search(r"valid to (\d{4}-\d{2}-\d{2})", note)
                if m and m.group(1) != c["valid_until_date"].isoformat():
                    w.append(f"{c['name'][:30]}: profile sheet says valid to {m.group(1)}, certificates sheet says {c['valid_until_date']} (using the certificates sheet)")
    for p in d["projects"]:
        if p["status"] == "completed" and not p["completion_date"]:
            w.append(f"{p['project_code']} is completed but has no completion date")
    return w


def _match_client(name: str, clients: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    best, score = None, 0
    for c in clients:
        s = fuzz.token_set_ratio(name.lower(), c["name"].lower())
        if s > score:
            best, score = c, s
    return best if score >= 70 else None


async def apply(d: Dict[str, Any], replace_demo: bool) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    async with AsyncSessionLocal() as db:
        if replace_demo:
            for model, col in ((Project, Project.project_code), (Employee, Employee.employee_code), (Certification, Certification.cert_code),
                               (Client, Client.client_code)):
                res = await db.execute(delete(model).where(col.is_(None)))
                counts[f"deleted demo {model.__tablename__}"] = res.rowcount

        async def upsert(model, key_col, key, values):
            row = (await db.execute(select(model).where(key_col == key))).scalar_one_or_none()
            if row is None:
                row = model(**{key_col.key: key})
                db.add(row)
            for k, v in values.items():
                setattr(row, k, v)
            return row

        for c in d["clients"]:
            await upsert(Client, Client.client_code, c["client_code"], {k: v for k, v in c.items() if k != "client_code"})
        await db.flush()
        by_name = [dict(id=c.id, name=c.name) for c in (await db.execute(select(Client))).scalars()]
        for p in d["projects"]:
            vals = {k: v for k, v in p.items() if k not in ("project_code", "client_name")}
            hit = _match_client(p["client_name"], by_name) if p["client_name"] else None
            if hit is None and p["client_name"]:
                cl = Client(name=p["client_name"], country="India")
                db.add(cl); await db.flush(); hit = dict(id=cl.id, name=cl.name); by_name.append(hit)
            vals["client_id"] = hit["id"] if hit else None
            await upsert(Project, Project.project_code, p["project_code"], vals)
        for e in d["employees"]:
            await upsert(Employee, Employee.employee_code, e["employee_code"], {k: v for k, v in e.items() if k != "employee_code"})
        for c in d["certifications"]:
            await upsert(Certification, Certification.cert_code, c["cert_code"], {k: v for k, v in c.items() if k != "cert_code"})
        for f in d["financials"]:
            await upsert(FinancialRecord, FinancialRecord.fiscal_year, f["fiscal_year"], {k: v for k, v in f.items() if k != "fiscal_year"} | {"currency": "INR"})
        for t in d["technologies"]:
            await upsert(Technology, Technology.name, t["name"], {k: v for k, v in t.items() if k != "name"})
        for key, (val, note, sensitive) in d["profile"].items():
            row = await db.get(CompanyProfile, key)
            if row is None:
                db.add(CompanyProfile(key=key, value=val, note=note, is_sensitive=sensitive))
            else:
                row.value, row.note, row.is_sensitive = val, note, sensitive
        await db.commit()
    return counts


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--replace-demo", action="store_true", help="delete rows that have no workbook ID (the demo seed data)")
    a = ap.parse_args()
    d = read_workbook(a.workbook)
    print("Read from workbook:", {k: len(v) for k, v in d.items() if k != "warnings"})
    for w in d["warnings"]:
        print("  WARNING:", w)
    if not a.apply:
        print("\nDry run only. Nothing written. Add --apply to import.")
        return
    extra = asyncio.run(apply(d, a.replace_demo))
    print("Imported.", extra or "")
    print("Updating the search index (needs the embedding model on Ollama)...")
    from app.services import kb_index

    async def _index():
        async with AsyncSessionLocal() as db:
            return await kb_index.rebuild(db)
    try:
        print("Index:", asyncio.run(_index()))
    except Exception as e:
        print(f"Index not updated ({type(e).__name__}); run it later from the Knowledge Base page or POST /kb/index/rebuild.")


if __name__ == "__main__":
    main()
