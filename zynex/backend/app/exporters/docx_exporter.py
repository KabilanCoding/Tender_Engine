import os
import uuid
from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_proposal_docx(proposal_data: dict, output_dir: str) -> str:
    """Generate a DOCX proposal document. Returns the file path."""
    doc = DocxDocument()

    # Title
    title = doc.add_heading(proposal_data.get("title", "Technical Proposal"), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Company name
    company_para = doc.add_paragraph()
    company_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = company_para.add_run(proposal_data.get("company_name", "ZYNEX Technologies"))
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()

    sections = proposal_data.get("sections", {})

    section_order = [
        ("executive_summary", "1. Executive Summary"),
        ("technical_approach", "2. Technical Approach"),
        ("project_plan", "3. Project Plan"),
        ("team_composition", "4. Team Composition"),
        ("company_profile", "5. Company Profile"),
        ("commercial_proposal", "6. Commercial Proposal"),
        ("compliance", "7. Compliance & Certifications"),
    ]

    for key, heading in section_order:
        content = sections.get(key) or proposal_data.get(key, "")
        if not content:
            continue
        doc.add_heading(heading, level=1)
        for para in content.split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.add_paragraph()

    # Save
    os.makedirs(output_dir, exist_ok=True)
    filename = f"proposal_{uuid.uuid4().hex[:8]}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath
